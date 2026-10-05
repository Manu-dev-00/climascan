import csv
import io
import json
from datetime import datetime
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import logging

logger = logging.getLogger(__name__)


class ExportService:
    """
    Génère les fichiers d'export au format CLIDATA :
    CSV normalisé WMO, Excel multi-feuilles, SQL Oracle.
    """

    CLIDATA_CSV_COLUMNS = [
        "STATION_ID", "ANNEE", "MOIS", "JOUR",
        "TN", "TX", "TMOY",
        "RR_TOTAL", "RR_DUREE",
        "ETP_PICHE", "ETP_BAC",
        "INSOLATION",
        "HR_0H", "HR_6H", "HR_12H", "HR_18H",
        "P_0H", "P_6H", "P_12H", "P_18H",
        "TENSION_VAPEUR",
        "VENT_MAX", "DIR_VENT",
        "CONFIDENCE_SCORE", "FLAGGED", "VALIDATED",
        "SCAN_DATE", "OPERATOR"
    ]

    def to_csv(self, records: List[Dict[str, Any]], include_metadata: bool = True) -> str:
        """Export CSV normalisé CLIDATA."""
        output = io.StringIO()
        writer = csv.DictWriter(
            output,
            fieldnames=self.CLIDATA_CSV_COLUMNS,
            extrasaction='ignore',
            delimiter=';'
        )
        writer.writeheader()

        for record in records:
            rows = self._flatten_record(record)
            for row in rows:
                if not include_metadata:
                    row.pop("CONFIDENCE_SCORE", None)
                    row.pop("FLAGGED", None)
                    row.pop("SCAN_DATE", None)
                    row.pop("OPERATOR", None)
                writer.writerow(row)

        return output.getvalue()

    def to_excel(self, records: List[Dict[str, Any]]) -> bytes:
        """Export Excel CLIDATA avec formatage professionnel."""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        # Couleurs
        header_fill = PatternFill("solid", fgColor="1B3A6B")
        flag_fill = PatternFill("solid", fgColor="FFE0E0")
        ok_fill = PatternFill("solid", fgColor="E8F5E9")
        alt_fill = PatternFill("solid", fgColor="F5F5F5")

        header_font = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
        body_font = Font(name="Calibri", size=10)
        bold_font = Font(name="Calibri", size=10, bold=True)
        thin = Side(style="thin", color="CCCCCC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for record in records:
            form_type = record.get("form_type", "donnees")[:20]
            mois = record.get("mois", "")
            annee = record.get("annee", "")
            sheet_name = f"{mois[:3]}_{annee}_{form_type[:10]}"
            ws = wb.create_sheet(title=sheet_name[:31])

            # Titre
            ws.merge_cells("A1:J1")
            ws["A1"] = f"Station DGMN Lomé — {mois} {annee} — {form_type}"
            ws["A1"].font = Font(bold=True, size=13, name="Calibri", color="1B3A6B")
            ws["A1"].alignment = Alignment(horizontal="center")
            ws.row_dimensions[1].height = 28

            # Infos station
            ws["A2"] = "Station :"
            ws["B2"] = record.get("station", "DGMN LOME")
            ws["C2"] = "Validé :"
            ws["D2"] = "Oui" if record.get("validated") else "Non"
            ws["E2"] = "Confiance globale :"
            ws["F2"] = f"{record.get('global_confidence', 0)*100:.1f}%"
            ws.row_dimensions[2].height = 18

            # En-têtes colonnes
            cols = self.CLIDATA_CSV_COLUMNS
            for ci, col in enumerate(cols, 1):
                cell = ws.cell(row=3, column=ci, value=col)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border
            ws.row_dimensions[3].height = 22

            # Données
            rows = self._flatten_record(record)
            for ri, row in enumerate(rows, 4):
                is_alt = (ri % 2 == 0)
                for ci, col in enumerate(cols, 1):
                    val = row.get(col, "")
                    cell = ws.cell(row=ri, column=ci, value=val)
                    cell.font = body_font
                    cell.border = border
                    cell.alignment = Alignment(horizontal="center")

                    # Coloration conditionnelle
                    if col == "FLAGGED" and val == "OUI":
                        ws.row_dimensions[ri].height = 16
                        for j in range(1, len(cols)+1):
                            ws.cell(row=ri, column=j).fill = flag_fill
                    elif col == "VALIDATED" and val == "OUI":
                        cell.fill = ok_fill
                    elif is_alt and col != "FLAGGED":
                        cell.fill = alt_fill

            # Largeurs colonnes
            for ci, col in enumerate(cols, 1):
                ws.column_dimensions[get_column_letter(ci)].width = max(
                    len(col) + 2, 10
                )

        # Feuille de métadonnées
        ws_meta = wb.create_sheet(title="Métadonnées", index=0)
        ws_meta["A1"] = "Export CliMaScan"
        ws_meta["A1"].font = Font(bold=True, size=14, color="1B3A6B")
        ws_meta["A2"] = f"Généré le : {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ws_meta["A3"] = f"Nombre de fiches : {len(records)}"
        ws_meta["A4"] = "Format : CLIDATA WMO"
        ws_meta["A5"] = "Station : DGMN Lomé — Latitude 06°10'N, Longitude 01°15'E"

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def to_sql(self, records: List[Dict[str, Any]]) -> str:
        """Génère les INSERT Oracle pour import CLIDATA direct."""
        lines = [
            "-- CliMaScan Export SQL Oracle — CLIDATA",
            f"-- Généré le {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            "-- Station : DGMN Lomé",
            "",
            "BEGIN",
            ""
        ]

        for record in records:
            rows = self._flatten_record(record)
            for row in rows:
                jour = row.get("JOUR", "NULL")
                mois = row.get("MOIS", "NULL")
                annee = row.get("ANNEE", "NULL")
                station = row.get("STATION_ID", "DGMN_LOME")

                def fmt(v):
                    if v is None or v == "":
                        return "NULL"
                    if isinstance(v, (int, float)):
                        return str(v)
                    return f"'{str(v).replace(chr(39), chr(39)+chr(39))}'"

                insert = (
                    f"  INSERT INTO CLIDATA_OBS "
                    f"(STATION_ID, ANNEE, MOIS, JOUR, TN, TX, TMOY, "
                    f"RR_TOTAL, ETP_PICHE, ETP_BAC, INSOLATION, "
                    f"CONFIDENCE_SCORE, FLAGGED, VALIDATED, SCAN_DATE) "
                    f"VALUES ("
                    f"{fmt(station)}, {fmt(annee)}, {fmt(mois)}, {fmt(jour)}, "
                    f"{fmt(row.get('TN'))}, {fmt(row.get('TX'))}, {fmt(row.get('TMOY'))}, "
                    f"{fmt(row.get('RR_TOTAL'))}, {fmt(row.get('ETP_PICHE'))}, "
                    f"{fmt(row.get('ETP_BAC'))}, {fmt(row.get('INSOLATION'))}, "
                    f"{fmt(row.get('CONFIDENCE_SCORE'))}, "
                    f"{fmt(1 if row.get('FLAGGED')=='OUI' else 0)}, "
                    f"{fmt(1 if row.get('VALIDATED')=='OUI' else 0)}, "
                    f"TO_DATE({fmt(row.get('SCAN_DATE',''))}, 'DD/MM/YYYY')"
                    f");"
                )
                lines.append(insert)

        lines += ["", "  COMMIT;", "END;", "/"]
        return "\n".join(lines)

    def _flatten_record(self, record: Dict) -> List[Dict]:
        """Aplatit un record structuré en lignes CSV/Excel."""
        rows = []
        daily = record.get("data", {}).get("daily", [])
        station = record.get("station", "DGMN_LOME")
        annee = record.get("annee", "")
        mois = record.get("mois", "")
        validated = "OUI" if record.get("validated") else "NON"
        scan_date = record.get("scan_at", datetime.now().strftime("%d/%m/%Y"))

        for day in daily:
            date = day.get("date", "")
            vals = day.get("values", {})

            def get_val(keys):
                for k in keys:
                    obj = vals.get(k)
                    if obj and isinstance(obj, dict):
                        return obj.get("value")
                return None

            def get_conf(keys):
                for k in keys:
                    obj = vals.get(k)
                    if obj and isinstance(obj, dict):
                        return obj.get("confidence", 1.0)
                return 1.0

            def is_flagged(keys):
                for k in keys:
                    obj = vals.get(k)
                    if obj and isinstance(obj, dict) and obj.get("flagged"):
                        return "OUI"
                return "NON"

            row = {
                "STATION_ID": station,
                "ANNEE": annee,
                "MOIS": mois,
                "JOUR": date,
                "TN": get_val(["Tn", "tn", "TN"]),
                "TX": get_val(["Tx", "tx", "TX"]),
                "TMOY": get_val(["Tmoy", "tmoy", "TMOY", "moyenne"]),
                "RR_TOTAL": get_val(["RR", "rr", "precip_total", "total"]),
                "RR_DUREE": get_val(["duree", "RR_DUREE"]),
                "ETP_PICHE": get_val(["piche_mm", "PICHE", "ETP_PICHE"]),
                "ETP_BAC": get_val(["bac_mm", "BAC", "ETP_BAC"]),
                "INSOLATION": get_val(["insolation", "INSOLATION"]),
                "HR_0H": get_val(["HR_0h", "hr_0h"]),
                "HR_6H": get_val(["HR_6h", "hr_6h"]),
                "HR_12H": get_val(["HR_12h", "hr_12h"]),
                "HR_18H": get_val(["HR_18h", "hr_18h"]),
                "P_0H": get_val(["P_0h", "p_0h"]),
                "P_6H": get_val(["P_6h", "p_6h"]),
                "P_12H": get_val(["P_12h", "p_12h"]),
                "P_18H": get_val(["P_18h", "p_18h"]),
                "TENSION_VAPEUR": get_val(["tension_vapeur", "TV"]),
                "VENT_MAX": get_val(["vent_max", "VENT_MAX"]),
                "DIR_VENT": get_val(["dir_vent", "DIR_VENT"]),
                "CONFIDENCE_SCORE": round(get_conf(list(vals.keys())[:3]), 3),
                "FLAGGED": is_flagged(list(vals.keys())),
                "VALIDATED": validated,
                "SCAN_DATE": scan_date if isinstance(scan_date, str)
                             else scan_date.strftime("%d/%m/%Y"),
                "OPERATOR": record.get("validated_by", ""),
            }
            rows.append(row)

        return rows


export_service = ExportService()
